"""Excel helpers: templates, reading uploads (xlsx/csv) and exporting tables."""

import csv
import io
from collections.abc import Iterable, Sequence
from datetime import date, datetime
from decimal import Decimal
from typing import Any
from uuid import UUID

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")

STUDENT_COLUMNS = [
    "admission_no",
    "full_name",
    "class",
    "section",
    "roll_no",
    "date_of_birth",
    "gender",
    "father_name",
    "father_phone",
    "mother_name",
    "mother_phone",
    "guardian_name",
    "guardian_phone",
]
TEACHER_COLUMNS = ["full_name", "phone", "employee_code", "email", "designation", "assignments", "class_teacher_of"]
MARKS_COLUMNS = ["admission_no", "roll_no", "student_name", "marks"]
META_SHEET = "_meta"


class SheetError(ValueError):
    pass


def _style_header(ws: Worksheet, columns: Sequence[str], widths: dict[str, int] | None = None) -> None:
    ws.append(list(columns))
    for i, col in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=i)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(col, max(14, len(col) + 4))
    ws.freeze_panes = "A2"


def _to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def students_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Students"
    _style_header(ws, STUDENT_COLUMNS, {"full_name": 26, "father_name": 22, "mother_name": 22})
    ws.append(
        [
            "A1001",
            "Example Student",
            "7",
            "A",
            "1",
            "2014-06-15",
            "F",
            "Example Father",
            "9876543210",
            "Example Mother",
            "9876543211",
            "",
            "",
        ]
    )
    help_ws = wb.create_sheet("Instructions")
    for line in [
        "One row per student. Required: admission_no, full_name, class, section.",
        "class and section must already exist for the current academic year (e.g. class 7, section A).",
        "date_of_birth: YYYY-MM-DD or DD/MM/YYYY. gender: M, F or O.",
        "Phones: 10-digit mobile numbers. Parents with the same phone are linked to all their children.",
        "Existing students (same admission_no) are updated, not duplicated.",
        "Delete the example row before uploading.",
    ]:
        help_ws.append([line])
    help_ws.column_dimensions["A"].width = 110
    return _to_bytes(wb)


def teachers_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Teachers"
    _style_header(ws, TEACHER_COLUMNS, {"full_name": 26, "assignments": 48, "email": 26})
    ws.append(["Example Teacher", "9876500001", "T001", "", "TGT Maths", "7-A:Mathematics; 7-B:Mathematics", "7-A"])
    help_ws = wb.create_sheet("Instructions")
    for line in [
        "Required: full_name, phone. Existing teachers (same phone) are updated.",
        "assignments: semicolon-separated 'class-section:Subject', e.g. 7-A:Mathematics; 8-B:Science",
        "class_teacher_of: the section this teacher is class teacher of, e.g. 7-A",
        "Delete the example row before uploading.",
    ]:
        help_ws.append([line])
    help_ws.column_dimensions["A"].width = 110
    return _to_bytes(wb)


def marks_template(
    *,
    exam_subject_id: UUID,
    section_id: UUID,
    title: str,
    max_marks: Decimal,
    students: Iterable[tuple[str, str | None, str]],
) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Marks"
    _style_header(ws, MARKS_COLUMNS, {"student_name": 30})
    for adm, roll, name in students:
        ws.append([adm, roll, name, None])
    note_col = len(MARKS_COLUMNS) + 2
    ws.cell(row=1, column=note_col, value=title).font = Font(bold=True)
    ws.cell(row=2, column=note_col, value=f"Max marks: {max_marks.normalize():f}")
    ws.cell(row=3, column=note_col, value="Enter a number, AB (absent) or EX (exempted).")
    ws.column_dimensions[get_column_letter(note_col)].width = 50
    meta = wb.create_sheet(META_SHEET)
    meta.append(["exam_subject_id", str(exam_subject_id)])
    meta.append(["section_id", str(section_id)])
    meta.sheet_state = "hidden"
    return _to_bytes(wb)


def _cell(v: Any) -> Any:
    if isinstance(v, str):
        v = v.strip()
        return v if v != "" else None
    return v


def read_rows(
    content: bytes, filename: str, columns: Sequence[str], max_rows: int
) -> tuple[list[tuple[int, dict[str, Any]]], dict[str, str]]:
    """Return ([(excel_row_number, {column: value})], meta) from an xlsx or csv upload."""
    name = filename.lower()
    meta: dict[str, str] = {}
    raw: list[list[Any]]
    if name.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        raw = [list(r) for r in csv.reader(io.StringIO(text))]
    elif name.endswith(".xlsx"):
        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise SheetError("File is not a valid .xlsx workbook") from exc
        ws = wb.worksheets[0]
        raw = [list(r) for r in ws.iter_rows(values_only=True)]
        if META_SHEET in wb.sheetnames:
            for m in wb[META_SHEET].iter_rows(values_only=True):
                if m and m[0] and len(m) > 1 and m[1] is not None:
                    meta[str(m[0])] = str(m[1])
        wb.close()
    else:
        raise SheetError("Upload an .xlsx or .csv file")
    if not raw:
        raise SheetError("The file is empty")
    header = [str(h).strip().lower() if h is not None else "" for h in raw[0]]
    missing = [c for c in columns if c not in header]
    if missing:
        raise SheetError(f"Missing columns: {', '.join(missing)}. Use the downloaded template.")
    idx = {c: header.index(c) for c in columns}
    rows: list[tuple[int, dict[str, Any]]] = []
    for n, r in enumerate(raw[1:], start=2):
        values = {c: _cell(r[i]) if i < len(r) else None for c, i in idx.items()}
        if all(v is None for v in values.values()):
            continue
        rows.append((n, values))
    if len(rows) > max_rows:
        raise SheetError(f"Too many rows ({len(rows)}). The limit is {max_rows} per file.")
    return rows, meta


def as_text(v: Any) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip() or None


def parse_date(v: Any) -> date | None:
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError("date must be YYYY-MM-DD or DD/MM/YYYY")


def export_table(title: str, columns: Sequence[tuple[str, str]], rows: Iterable[dict[str, Any]]) -> bytes:
    """columns: [(key, header)]."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = title[:31]
    _style_header(ws, [h for _, h in columns])
    for r in rows:
        ws.append([_export_value(r.get(k)) for k, _ in columns])
    return _to_bytes(wb)


def _export_value(v: Any) -> Any:
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, UUID):
        return str(v)
    if isinstance(v, list):
        return "; ".join(str(x) for x in v)
    if hasattr(v, "value"):
        return v.value
    return v
