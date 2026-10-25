"""Dashboard numbers on a hand-checked fixture.

Unit Test 1 (max 100 each; Maths, Science, English):
  7-A  S1: 30, 20, 90  -> 46.67%, 2 fails
       S2: 90, 80, 70  -> 80.00%
       S3: 60, 60, 60  -> 60.00%
  7-B  S4: 50, -, -    (only Maths applies? no: all 3 subjects; S4 gets 50,50,50 -> 50.00%)
Half-Yearly: S2 drops to 65.00% (-15) -> needs attention; S3 stays 60.
"""

import io

import pytest
from openpyxl import load_workbook

from tests.factories import put_marks


@pytest.fixture
def dash(client, world):  # type: ignore[no-untyped-def]
    a = world.a
    h = a.h("admin")
    data = {a.s1: (30, 20, 90), a.s2: (90, 80, 70), a.s3: (60, 60, 60)}
    for i, es in enumerate((a.e1_math, a.e1_sci, a.e1_eng)):
        put_marks(client, h, es, a.sec_a, [{"student_id": s, "marks": v[i]} for s, v in data.items()])
        put_marks(client, h, es, a.sec_b, [{"student_id": a.s4, "marks": 50}])
    for es in (a.e2_math, a.e2_sci, a.e2_eng):
        put_marks(client, h, es, a.sec_a, [{"student_id": a.s2, "marks": 65}, {"student_id": a.s3, "marks": 60}])
    return world


def _q(client, w, path):  # type: ignore[no-untyped-def]
    r = client.get(f"/api/v1/dashboard/{path}", headers=w.a.h("principal"))
    assert r.status_code == 200, r.text
    return r.json()


def test_summary(client, dash) -> None:  # type: ignore[no-untyped-def]
    s = _q(client, dash, f"summary?exam_id={dash.a.e1}")
    assert s["students"] == 4 and s["students_with_results"] == 4
    # (46.67 + 80 + 60 + 50) / 4 = 59.1675 -> 59.17
    assert s["average_percent"] == "59.17"
    assert s["pass_percent"] == "75.00"  # S1 failed two subjects
    dist = {g["grade"]: g["count"] for g in s["grade_distribution"]}
    assert dist == {"A1": 0, "A2": 0, "B1": 1, "B2": 0, "C1": 1, "C2": 2, "D": 0, "E": 0}
    assert s["sections_total"] == 2 and s["sections_published"] == 0


def test_sections_and_subjects(client, dash) -> None:  # type: ignore[no-untyped-def]
    secs = {s["section_name"]: s for s in _q(client, dash, f"sections?exam_id={dash.a.e1}")}
    assert secs["Class 7-A"]["average_percent"] == "62.22" and secs["Class 7-A"]["students"] == 3
    assert secs["Class 7-B"]["average_percent"] == "50.00"
    subs = _q(client, dash, f"subjects?exam_id={dash.a.e1}")
    # Maths (30+90+60+50)/4=57.5, Science (20+80+60+50)/4=52.5, English (90+70+60+50)/4=67.5
    assert [(s["subject_name"], s["average_percent"]) for s in subs] == [
        ("Science", "52.50"),
        ("Mathematics", "57.50"),
        ("English", "67.50"),
    ]
    assert subs[0]["pass_percent"] == "75.00"  # S1 scored 20 < 33
    only_a = _q(client, dash, f"subjects?exam_id={dash.a.e1}&section_id={dash.a.sec_a}")
    assert next(s for s in only_a if s["subject_name"] == "Mathematics")["average_percent"] == "60.00"


def test_attention_list(client, dash) -> None:  # type: ignore[no-untyped-def]
    a = dash.a
    ut1 = _q(client, dash, f"attention?exam_id={a.e1}")
    assert [(x["full_name"], x["subjects_failed"]) for x in ut1] == [("Asha SCHA", 2)]
    hy = _q(client, dash, f"attention?exam_id={a.e2}")
    assert [(x["full_name"], x["drop"], x["previous_percentage"]) for x in hy] == [("Bala SCHA", "15.00", "80.00")]
    assert "Dropped 15.00 points" in hy[0]["reasons"][0]


def test_toppers_and_drilldown(client, dash) -> None:  # type: ignore[no-untyped-def]
    a = dash.a
    tops = _q(client, dash, f"toppers?exam_id={a.e1}&top=2")
    assert [(t["section_name"], t["rank"], t["full_name"]) for t in tops] == [
        ("Class 7-A", 1, "Bala SCHA"),
        ("Class 7-A", 2, "Chitra SCHA"),
        ("Class 7-B", 1, "Deepak SCHA"),
    ]
    rows = _q(client, dash, f"students?exam_id={a.e1}&section_id={a.sec_a}")
    assert [(r["full_name"], r["rank"], r["grade"]) for r in rows] == [
        ("Bala SCHA", 1, "B1"),
        ("Chitra SCHA", 2, "C1"),
        ("Asha SCHA", 3, "C2"),
    ]


def test_export_xlsx_and_audit(client, dash) -> None:  # type: ignore[no-untyped-def]
    a = dash.a
    for report in ("summary", "sections", "subjects", "attention", "toppers"):
        r = client.get(f"/api/v1/dashboard/export?report={report}&exam_id={a.e1}", headers=a.h("principal"))
        assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats"), report
        assert load_workbook(io.BytesIO(r.content)).worksheets[0].max_row >= 2
    r = client.get(
        f"/api/v1/dashboard/export?report=students&exam_id={a.e1}&section_id={a.sec_a}", headers=a.h("principal")
    )
    ws = load_workbook(io.BytesIO(r.content)).worksheets[0]
    assert ws.max_row == 4 and ws["C2"].value == "Bala SCHA"
    logs = client.get("/api/v1/audit-logs", headers=a.h("principal")).json()
    assert sum(1 for x in logs if x["action"] == "EXPORT") == 6


def test_dashboard_roles(client, dash) -> None:  # type: ignore[no-untyped-def]
    a = dash.a
    for who in ("teacher1", "parent1", "student_s1"):
        r = client.get(f"/api/v1/dashboard/summary?exam_id={a.e1}", headers=a.h(who))
        assert r.status_code == 403
    assert client.get(f"/api/v1/dashboard/summary?exam_id={a.e1}", headers=a.h("admin")).status_code == 200


def test_dashboard_performance_1500_students(client, world, owner_engine) -> None:  # type: ignore[no-untyped-def]
    """MVP_SPEC §9: dashboard API < 500 ms for a 1,500-student school (bulk-loaded fixture)."""
    import time

    from sqlalchemy import text

    a = world.a
    with owner_engine.begin() as conn:
        conn.execute(
            text("""
            WITH s AS (
              INSERT INTO students (id, school_id, admission_no, full_name)
              SELECT gen_random_uuid(), :sch, 'P' || g, 'Perf ' || g FROM generate_series(1, 1500) g
              RETURNING id, admission_no)
            INSERT INTO student_enrolments (id, school_id, student_id, academic_year_id, section_id)
            SELECT gen_random_uuid(), :sch, s.id, :yr,
                   CASE WHEN substr(s.admission_no, 2)::int % 2 = 0 THEN :sa ELSE :sb END FROM s"""),
            {"sch": a.school_id, "yr": a.year_id, "sa": a.sec_a, "sb": a.sec_b},
        )
        for es in (a.e1_math, a.e1_sci, a.e1_eng, a.e2_math, a.e2_sci, a.e2_eng):
            conn.execute(
                text("""
                INSERT INTO marks (id, school_id, exam_subject_id, student_id, section_id, status, marks_obtained)
                SELECT gen_random_uuid(), :sch, CAST(:es AS uuid), e.student_id, e.section_id, 'PRESENT',
                       (abs(hashtext(e.student_id::text || CAST(:es AS text))) % 101)
                FROM student_enrolments e WHERE e.school_id = :sch"""),
                {"sch": a.school_id, "es": str(es)},
            )
    h = a.h("principal")
    for path in ("summary", "sections", "subjects", "attention", "toppers"):
        client.get(f"/api/v1/dashboard/{path}?exam_id={a.e2}", headers=h)  # warm up
        start = time.perf_counter()
        r = client.get(f"/api/v1/dashboard/{path}?exam_id={a.e2}", headers=h)
        elapsed = (time.perf_counter() - start) * 1000
        assert r.status_code == 200
        assert elapsed < 500, (path, elapsed)
