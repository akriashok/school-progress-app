import io
import time

from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.core.db import SessionLocal
from app.models import AuditLog
from app.services.excel import MARKS_COLUMNS, STUDENT_COLUMNS, TEACHER_COLUMNS


def _xlsx(columns: list[str], rows: list[list[object]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(columns)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(client, h, kind: str, content: bytes, name: str = "f.xlsx", params: str = ""):  # type: ignore[no-untyped-def]
    r = client.post(f"/api/v1/imports/{kind}{params}", headers=h, files={"file": (name, content)})
    assert r.status_code == 202, r.text
    return client.get(f"/api/v1/imports/{r.json()['id']}", headers=h).json()  # background job has run


def test_templates_download(client, world) -> None:  # type: ignore[no-untyped-def]
    h = world.a.h("admin")
    for kind, cols in (("students", STUDENT_COLUMNS), ("teachers", TEACHER_COLUMNS)):
        r = client.get(f"/api/v1/imports/templates/{kind}", headers=h)
        assert r.status_code == 200
        ws = load_workbook(io.BytesIO(r.content)).worksheets[0]
        assert [c.value for c in ws[1]] == cols
    assert client.get("/api/v1/imports/templates/students", headers=world.a.h("teacher1")).status_code == 403


def test_student_import_preview_errors_and_commit(client, world) -> None:  # type: ignore[no-untyped-def]
    h = world.a.h("admin")
    rows = [
        ["N100", "Ravi Teja", "7", "A", "4", "2014-05-01", "M", "Ravi Sr", "9844444444", "", "", "", ""],
        ["N101", "Sita", "Class 7", "b", "", "01/02/2014", "F", "", "", "Sita Mom", "9844444445", "", ""],
        ["N102", "No Section", "9", "Z", "", "", "", "", "", "", "", "", ""],
        ["N103", "Bad Phone", "7", "A", "", "", "", "Dad", "12345", "", "", "", ""],
        ["N100", "Duplicate", "7", "A", "", "", "", "", "", "", "", "", ""],
        ["", "", "", "", "", "", "", "", "", "", "", "", ""],  # blank row ignored
        ["N104", "Bad Date", "7", "A", "", "31-31-2014", "X", "", "", "", "", "", ""],
    ]
    imp = _upload(client, h, "students", _xlsx(STUDENT_COLUMNS, rows))
    assert imp["status"] == "READY" and imp["total_rows"] == 6 and imp["valid_rows"] == 2
    msgs = {(e["row"], e["column"]): e["message"] for e in imp["errors"]}
    assert msgs[(4, "section")] == "section 9-Z does not exist in the current year"
    assert msgs[(5, "father_phone")] == "phone must be 10 digits"
    assert msgs[(6, "admission_no")] == "duplicate of row 2"
    assert (8, "date_of_birth") in msgs and (8, "gender") in msgs
    assert [p["admission_no"] for p in imp["preview"]] == ["N100", "N101"]

    errs = client.get(f"/api/v1/imports/{imp['id']}/errors.xlsx", headers=h)
    assert errs.status_code == 200 and load_workbook(io.BytesIO(errs.content)).worksheets[0].max_row == 6

    r = client.post(f"/api/v1/imports/{imp['id']}/commit", headers=h, json={})
    assert r.status_code == 200 and r.json()["status"] == "COMMITTED"
    assert r.json()["result"] == {"created": 2, "updated": 0, "guardian_links": 2}
    assert client.post(f"/api/v1/imports/{imp['id']}/commit", headers=h, json={}).status_code == 409
    found = client.get("/api/v1/students?q=Sita", headers=h).json()["items"]
    assert found[0]["section_name"] == "Class 7-B" and found[0]["date_of_birth"] == "2014-02-01"
    # Parent created and can log in.
    v = client.post("/api/v1/auth/otp/verify", json={"phone": "9844444445", "otp": "123456"}).json()
    assert [x["role"] for x in v["roles"]] == ["PARENT"]
    with SessionLocal() as db:
        assert db.scalar(select(AuditLog).where(AuditLog.action == "IMPORT_COMMIT")) is not None

    # Re-importing upserts by admission_no instead of duplicating.
    imp2 = _upload(client, h, "students", _xlsx(STUDENT_COLUMNS, [rows[0][:1] + ["Ravi Teja K"] + rows[0][2:]]))
    r = client.post(f"/api/v1/imports/{imp2['id']}/commit", headers=h, json={})
    assert r.json()["result"]["updated"] == 1 and r.json()["result"]["created"] == 0


def test_teacher_import(client, world) -> None:  # type: ignore[no-untyped-def]
    a = world.a
    h = a.h("admin")
    rows = [
        ["Kiran", "9855555555", "T9", "k@example.com", "PGT", "7-A:Mathematics; 7-b:SCI", "7-B"],
        ["Nope", "9855555556", "", "", "", "9-Z:Mathematics", ""],
        ["Nope2", "9855555557", "", "", "", "7-A:Dance", ""],
    ]
    imp = _upload(client, h, "teachers", _xlsx(TEACHER_COLUMNS, rows))
    assert imp["valid_rows"] == 1
    assert [e["message"] for e in imp["errors"]] == ["unknown section '9-Z'", "unknown subject 'Dance'"]
    r = client.post(f"/api/v1/imports/{imp['id']}/commit", headers=h, json={})
    assert r.json()["result"] == {"created": 1, "updated": 0, "assignments_added": 2}
    sec_b = client.get(f"/api/v1/sections/{a.sec_b}", headers=h).json()
    assert sec_b["class_teacher_name"] == "Kiran"


def test_marks_template_round_trip(client, world) -> None:  # type: ignore[no-untyped-def]
    a = world.a
    h = a.h("teacher1")
    params = f"?exam_subject_id={a.e1_math}&section_id={a.sec_a}"
    t = client.get(f"/api/v1/imports/templates/marks{params}", headers=h)
    assert t.status_code == 200
    wb = load_workbook(io.BytesIO(t.content))
    ws = wb["Marks"]
    assert [ws.cell(row=i, column=3).value for i in (2, 3, 4)] == ["Asha SCHA", "Bala SCHA", "Chitra SCHA"]
    assert wb["_meta"].sheet_state == "hidden"
    ws.cell(row=2, column=4, value=88)
    ws.cell(row=3, column=4, value="AB")
    ws.cell(row=4, column=4, value=120)
    buf = io.BytesIO()
    wb.save(buf)
    imp = _upload(client, h, "marks", buf.getvalue(), params=params)
    assert imp["valid_rows"] == 2
    assert imp["errors"] == [{"row": 4, "column": "marks", "message": "Row 4: marks 120 > max 100"}]
    r = client.post(f"/api/v1/imports/{imp['id']}/commit", headers=h, json={})
    assert r.json()["result"] == {"created": 2, "updated": 0, "unchanged": 0}
    grid = client.get(f"/api/v1/marks{params}", headers=h).json()
    assert [(x["status"], x["marks"]) for x in grid["rows"]] == [("PRESENT", "88.00"), ("ABSENT", None), (None, None)]


def test_marks_import_permissions_and_wrong_file(client, world) -> None:  # type: ignore[no-untyped-def]
    a = world.a
    params = f"?exam_subject_id={a.e1_math}&section_id={a.sec_a}"
    content = _xlsx(MARKS_COLUMNS, [["SCHA-S1", "1", "Asha", 10]])
    r = client.post(f"/api/v1/imports/marks{params}", headers=a.h("teacher2"), files={"file": ("m.xlsx", content)})
    assert r.status_code == 403
    for who in ("parent1", "student_s1"):
        r = client.post(f"/api/v1/imports/marks{params}", headers=a.h(who), files={"file": ("m.xlsx", content)})
        assert r.status_code == 403
    # Template for 7-B uploaded against 7-A
    other = client.get(
        f"/api/v1/imports/templates/marks?exam_subject_id={a.e1_math}&section_id={a.sec_b}", headers=a.h("admin")
    ).content
    imp = _upload(client, a.h("admin"), "marks", other, params=params)
    assert imp["valid_rows"] == 0 and "different exam" in imp["errors"][0]["message"]
    # Teachers can't upload student lists.
    r = client.post("/api/v1/imports/students", headers=a.h("teacher1"), files={"file": ("s.xlsx", content)})
    assert r.status_code == 403
    # Another school can't read the import.
    assert client.get(f"/api/v1/imports/{imp['id']}", headers=world.b.h("admin")).status_code == 404


def test_bad_files(client, world) -> None:  # type: ignore[no-untyped-def]
    h = world.a.h("admin")
    r = client.post("/api/v1/imports/students", headers=h, files={"file": ("x.pdf", b"%PDF-1.4")})
    assert r.status_code == 422 and r.json()["error"]["code"] == "FILE_TYPE"
    r = client.post("/api/v1/imports/students", headers=h, files={"file": ("x.xlsx", b"not a zip")})
    assert r.status_code == 422
    imp = _upload(client, h, "students", _xlsx(["foo"], [["bar"]]))
    assert imp["status"] == "FAILED" and "Missing columns" in imp["errors"][0]["message"]
    csv_body = ",".join(STUDENT_COLUMNS) + "\nC1,Csv Kid,7,A,,,,,,,,,\n"
    imp = _upload(client, h, "students", csv_body.encode(), name="s.csv")
    assert imp["status"] == "READY" and imp["valid_rows"] == 1


def test_500_row_import_is_fast(client, world) -> None:  # type: ignore[no-untyped-def]
    h = world.a.h("admin")
    rows = [
        [f"B{i}", f"Kid {i}", "7", "A" if i % 2 else "B", str(i), "", "F", f"Dad {i}", f"97{i:08d}", "", "", "", ""]
        for i in range(500)
    ]
    start = time.perf_counter()
    imp = _upload(client, h, "students", _xlsx(STUDENT_COLUMNS, rows))
    r = client.post(f"/api/v1/imports/{imp['id']}/commit", headers=h, json={})
    elapsed = time.perf_counter() - start
    assert r.json()["result"]["created"] == 500
    assert elapsed < 30, elapsed
